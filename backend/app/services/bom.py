"""BOM 文本解析与取料规划（纯函数层，无 IO，可独立单测）。

职责（与前端 BomDialog / GuideOverlay 对应）：
1. parse_text：把粘贴的 BOM 文本按分隔符拆行，逐行还原出
   name / value / package / quantity 四要素，数量缺省 1；
2. normalize_value：把「0.1uF / 100nF / 10k / 22Ω」归一化成同一
   标度 + 家族，供匹配做等值判断（相对容差 1e-9）；
3. plan_pure：把解析结果与库存对照，产出「引导步骤 + 缺料清单」，
   同一元件多行合并为一步取总量，步骤按 led_index 升序（None 排最后）。

解析顺序约定：先剥离数量标记（x20 / ×5 / *3 / 20个 / 10 只 / 5pcs），
再识别值/封装 token——避免把「10k」里的 10 当数量。一行出现多个数量
标记时取最靠前的一个（其余剥离）。四位纯数字留给封装识别，1~3 位裸
数字视为无后缀阻值。

Excel 导入（parse_excel_bytes）：识别常见表头（名称/规格/封装/数量/
Manufacturer Part/Supplier Part），分别保留字段，避免把型号中的数字
当作数量或封装；料号用于采购与识别。
"""

import csv
import json
import re
import unicodedata
from dataclasses import dataclass
from difflib import SequenceMatcher
from io import BytesIO, StringIO
from typing import Any, Optional

from openpyxl import load_workbook

# ---------- 常量：token 形态 ----------

_NUM = r"(?:\d+(?:\.\d+)?|\.\d+)"
# 单位前缀（p/n/u/μ/m/k/M/G，注意 m 与 M 语义不同不可合写）
_PREFIX = r"(?:p|n|u|U|μ|m|k|K|M|G)"
# 家族后缀：F 电容 / R、Ω、欧、ohm 电阻 / H 电感
_SUFFIX = r"(?:F|f|R|r|Ω|欧|ohm|OHM|Ohm|H|h)"

# 数量形态：x20 / ×5 / *3 / 20个 / 10 只 / 5pcs（re.IGNORECASE 覆盖 pcs）
_QTY_RE = re.compile(
    r"(?:[x×*]\s*(\d+))|(?:(\d+)\s*(?:个|只|颗|枚|片|pcs))",
    re.IGNORECASE,
)
# 值 token：带单位前缀 → 直接带家族后缀 → 1~3 位裸数字（四位留给封装）
_VALUE_RE = re.compile(
    r"(?<![A-Za-z0-9_.-])"
    r"(?:"
    r"\d+[RrKkM]\d+(?:Ω|欧)?"
    r"|{num}\s*{prefix}\s*{suffix}?"
    r"|{num}\s*{suffix}"
    r"|\d{{1,3}}(?![A-Za-z0-9.])"
    r")"
    r"(?![A-Za-z0-9_.-])".format(num=_NUM, prefix=_PREFIX, suffix=_SUFFIX)
)
# 封装 token：独立四位数字（0603 / 0805 / 1206 …）
_PACK_RE = re.compile(
    r"(?<![A-Za-z0-9_.-])(?:\d{4}|"
    r"(?:SOT|SOIC|SOP|SSOP|TSSOP|QFN|LQFP|DIP|TO)[ -]?\d{1,3})"
    r"(?![A-Za-z0-9_.-])", re.IGNORECASE,
)

# 归一化正则（组件库里 value 字段的宽容解析，允许 "22" 裸数字）
_NORM_RE = re.compile(
    r"^\s*(?P<num>\d+(?:\.\d+)?|\.\d+)\s*"
    r"(?P<unit>p|n|u|U|μ|m|k|K|M|G)?\s*"
    r"(?P<fam>F|f|R|r|Ω|欧|ohm|OHM|Ohm|H|h)?\s*$"
)
_MULT = {"p": 1e-12, "n": 1e-9, "u": 1e-6, "U": 1e-6, "μ": 1e-6,
         "m": 1e-3, "k": 1e3, "K": 1e3, "M": 1e6, "G": 1e9}
_FAMILY = {"F": "F", "f": "F", "R": "R", "r": "R", "Ω": "R",
           "欧": "R", "ohm": "R", "OHM": "R", "Ohm": "R",
           "H": "H", "h": "H"}

# 名称碎片：数量/单位残留与孤标点（如 "x12"、"uF"、"。"）
_LEFTOVER_RE = re.compile(
    r"^[x×*]\d*$|^[pnuμUmMkKGFRrΩ欧Hho]+$|^[.。，,、:：;；()（）]+$"
)


@dataclass
class BomLine:
    """BOM 文本中一行解析出的元件诉求（Excel 行还会带厂商料号等）。"""

    raw: str
    name: str = ""
    value: Optional[str] = None
    package: Optional[str] = None
    quantity: int = 1
    manufacturer_part: Optional[str] = None
    supplier_part: Optional[str] = None

    def to_dict(self) -> dict:
        """转 JSON 字典（/bom/parse 与 /bom/plan 响应复用）。

        文本行没有料号，省略键保持向后兼容；Excel 行才带。
        """
        out = {
            "raw": self.raw,
            "name": self.name,
            "value": self.value,
            "package": self.package,
            "quantity": self.quantity,
        }
        if self.manufacturer_part:
            out["manufacturer_part"] = self.manufacturer_part
        if self.supplier_part:
            out["supplier_part"] = self.supplier_part
        return out


def _blank_all(text: str, pattern: re.Pattern) -> tuple[Optional[str], str]:
    """把所有匹配替换为空格（防 token 粘连），返回（首个匹配串, 清理后文本）。"""
    matches = list(pattern.finditer(text))
    first = matches[0].group().strip() if matches else None
    # 倒序替换保证下标不因长度变化而错位
    for match in reversed(matches):
        text = text[: match.start()] + " " + text[match.end():]
    return first, text


def _clean_name(work: str) -> str:
    """把 token 里剥离后残留的数字/单位碎片清掉，拼出纯名称。"""
    kept = []
    for token in work.split():
        if not token or token.isdigit() or _LEFTOVER_RE.fullmatch(token):
            continue
        kept.append(token)
    return " ".join(kept).strip()


def parse_text(text: str) -> list[BomLine]:
    """把 BOM 文本按 [,，;；\n]+ 拆行，逐行解析成 BomLine 列表。"""
    lines: list[BomLine] = []
    for raw in re.split(r"[,，;；\n\r]+", text):
        raw = raw.strip()
        if not raw:
            continue

        # 1) 先剥离全部数量标记；多个标记取最靠前的一个
        quantity = 1
        seen = False
        for match in _QTY_RE.finditer(raw):
            if seen:
                break
            quantity = int(match.group(1) or match.group(2))
            seen = True
        work = _QTY_RE.sub(" ", raw)

        # 2) 值 token：首个进 value，其余全部剥离
        value_text, work = _blank_all(work, _VALUE_RE)
        if value_text:
            value_text = re.sub(r"\s+", "", value_text)

        # 3) 封装 token：四位独立数字
        package, work = _blank_all(work, _PACK_RE)

        # 4) 剩余内容做碎片清理后即名称
        name = _clean_name(work)

        # 一行什么都没留下（如只有数量/纯标点）视为不可识别，跳过
        if not name and not value_text and not package:
            continue
        lines.append(BomLine(raw=raw, name=name, value=value_text,
                             package=package, quantity=quantity))
    return lines


def normalize_value(text: str) -> Optional[tuple[float, str]]:
    """把值文本换算成 (标度化数值, 家族)，无法识别返回 None。

    例：10k → (10000.0, R)；100nF 与 0.1uF 都归到 1e-07 F；
    空后缀（22 / 4.7k）默认电阻族。
    """
    text = unicodedata.normalize("NFKC", text).strip()
    text = re.sub(r"(?i)ohms?|欧姆", "Ω", text)
    text = re.sub(
        r"^(\d+)([RrKkM])(\d+)(Ω|欧)?$",
        lambda m: f"{m[1]}.{m[3]}{m[2] if m[2] not in 'Rr' else ''}Ω",
        text,
    )
    match = _NORM_RE.fullmatch(text)
    if match is None:
        return None
    num = float(match.group("num"))
    unit = match.group("unit")
    fam_raw = match.group("fam")
    scaled = num * (_MULT.get(unit, 1.0) if unit else 1.0)
    family = _FAMILY.get(fam_raw, "R") if fam_raw else "R"
    return scaled, family


def norm_values_equal(left: str, right: str) -> bool:
    """两个值文本是否等价：家族一致 + 数值相对容差 1e-9。"""
    a = normalize_value(left)
    b = normalize_value(right)
    if a is None or b is None:
        return False
    (x, fa), (y, fb) = a, b
    if fa != fb:
        return False
    return x == y or abs(x - y) <= 1e-9 * max(abs(x), abs(y))


_NAME_ALIASES = (
    ("resistor", "电阻"), ("capacitor", "电容"), ("inductor", "电感"),
    ("发光二极管", "led"), ("指示灯", "led"), ("贴片", ""), ("插件", ""),
)
_GENERIC_NAMES = {
    "电阻", "电容", "电感", "led", "芯片", "元件", "连接器", "二极管",
    "电解电容", "陶瓷电容", "稳压器", "电源芯片", "testpoint",
}


def _compact(text: Optional[str]) -> str:
    """人类字段归一化：大小写、全角、空格与分隔符。"""
    return "".join(c for c in unicodedata.normalize(
        "NFKC", text or "",
    ).casefold() if c.isalnum() or c == ".")


def _name_key(text: str) -> str:
    """统一常见类别别名，保留型号数字与后缀。"""
    text = _compact(text)
    for old, new in _NAME_ALIASES:
        text = text.replace(old, new)
    return text


def _package_key(text: Optional[str]) -> str:
    """统一 R0603/0603、SOT-223/SOT223 等写法。"""
    key = _compact(text)
    match = re.fullmatch(r"[a-z]{1,3}(\d{4})", key)
    return match[1] if match else key


def _spec_value(text: Optional[str]) -> Optional[tuple[float, str]]:
    """读取完整标称值或描述中的唯一标称值，避免混合规格误推断。"""
    if not text:
        return None
    direct = normalize_value(text)
    if direct is not None:
        return direct
    text = unicodedata.normalize("NFKC", text)
    values = {normalize_value(m[0]) for m in _VALUE_RE.finditer(text)
              if not m[0].strip().isdigit()}
    values.discard(None)
    return next(iter(values)) if len(values) == 1 else None


def _name_similarity(left: str, right: str) -> int:
    """名称精确、包含或近似评分；短碎片不参与包含匹配。"""
    if not left or not right:
        return 0
    if left == right:
        return 60
    if min(len(left), len(right)) >= 2 and (left in right or right in left):
        return 32
    ratio = SequenceMatcher(None, left, right).ratio()
    return int(20 * ratio) if ratio >= 0.7 else 0


def _score_line(line: BomLine, comp: Any) -> tuple[int, list[str], bool]:
    """返回评分、依据及自动匹配资格；冲突的料号命中只供人工确认。"""
    tags = getattr(comp, "tags", None) or []
    if isinstance(tags, str):
        try:
            tags = json.loads(tags)
        except ValueError:
            tags = []
    tags = [str(tag) for tag in tags] if isinstance(tags, list) else []
    pairs = [
        (line.supplier_part, getattr(comp, "supplier_part", None),
         1000, "供应商料号相同"),
        (line.manufacturer_part, getattr(comp, "manufacturer_part", None),
         500, "厂商料号相同"),
    ]
    part_matches = [p for p in pairs if p[0] and p[1]
                    and p[0].strip().casefold() == p[1].strip().casefold()]
    part_conflict = any(p[0] and p[1] and p not in part_matches for p in pairs)
    if part_conflict and not part_matches:
        return 0, [], False
    reasons = [p[3] for p in part_matches]
    score = sum(p[2] for p in part_matches)
    conflicts = ["料号不一致"] if part_conflict else []
    name = _name_key(line.name)
    comp_name = _name_key(comp.name or "")
    name_score = _name_similarity(name, comp_name)
    tag_score = max((_name_similarity(name, _name_key(t)) for t in tags),
                    default=0)
    lp = _package_key(line.package)
    cp = _package_key(comp.package)
    if not cp:
        cp = next((_package_key(t) for t in tags if _package_key(t) == lp), "")
    package_equal = bool(lp and cp and lp == cp)
    if lp and cp and lp != cp:
        conflicts.append("封装不一致")
    lv = _spec_value(line.value)
    cv = _spec_value(comp.value)
    if cv is None and not comp.value:
        inferred = {_spec_value(t) for t in [comp.name or "", *tags]
                    if not t.strip().isdigit()}
        inferred.discard(None)
        cv = next(iter(inferred)) if len(inferred) == 1 else None
    value_equal = bool(lv is not None and cv is not None
                       and lv[1] == cv[1] and abs(lv[0] - cv[0])
                       <= 1e-9 * max(abs(lv[0]), abs(cv[0])))
    if lv is not None and cv is not None and not value_equal:
        conflicts.append("标称值不一致")
    elif line.value and comp.value and (lv is None or cv is None):
        value_equal = _compact(line.value) == _compact(comp.value)
        if not value_equal:
            conflicts.append("标称值不一致")
    if any(kind in name and other in comp_name
           for kind in ("电阻", "电容", "电感")
           for other in ("电阻", "电容", "电感") if kind != other):
        conflicts.append("元件类别不一致")
    if conflicts and not part_matches:
        return 0, [], False
    if not (part_matches or value_equal or name_score or tag_score
            or (package_equal and not line.name and not line.value)):
        return 0, [], False
    if value_equal:
        reasons.append("标称值相同")
        score += 40
    if package_equal:
        reasons.append("封装相同")
        score += 25
    if name_score:
        reasons.append("名称相同" if name_score == 60 else "名称相近")
        score += name_score
    if tag_score > name_score:
        reasons.append("标签命中" if tag_score == 60 else "标签相近")
        score += tag_score
    automatic = bool(part_matches or (value_equal and lv is not None and
                     (package_equal or name_score >= 32) and
                     (not name or name in _GENERIC_NAMES or
                      name_score >= 32 or tag_score >= 32)) or
                     (name not in _GENERIC_NAMES and
                      (name_score == 60 or tag_score == 60)))
    return score, reasons + conflicts, automatic and not conflicts


def plan_pure(
    lines: list[BomLine], components: list[Any],
    selections: Optional[dict[int, int]] = None,
) -> dict[str, Any]:
    """库存对照规划（纯函数）。

    返回：
      requested     全部行数量之和
      steps         [{component, quantity, line_indexes}]，足料且按
                    led_index 升序（None 排最后）；同一元件多行合并取总量
      missing       [{reason, raw, name, value, package, quantity,
                      component_id?, available?}]，not_found 每行一条、
                    shortage 按元件合并一条
      choices       [{line_index, line, candidates}]，近似或歧义结果待确认
    """
    selections = selections or {}
    if any(index < 0 or index >= len(lines) for index in selections):
        raise ValueError("BOM 行号无效")
    matched: dict[int, Any] = {}
    choices: list[dict] = []
    for index, line in enumerate(lines):
        ranked = []
        for comp in components:
            score, reasons, automatic = _score_line(line, comp)
            if score > 0:
                ranked.append({"component": comp, "score": score,
                               "reasons": reasons, "automatic": automatic})
        ranked.sort(key=lambda item: item["score"], reverse=True)
        ranked = ranked[:5]
        if index in selections:
            chosen = next((item for item in ranked
                           if item["component"].id == selections[index]), None)
            if chosen is None:
                raise ValueError("所选元件不在当前候选中")
            matched[index] = chosen["component"]
        elif ranked:
            best = ranked[0]
            close = [item for item in ranked
                     if item["score"] >= best["score"] - 15]
            if best["automatic"] and len(close) == 1:
                matched[index] = best["component"]
            else:
                choices.append({"line_index": index, "line": line.to_dict(),
                                "candidates": ranked})

    # 同一元件多行合并：need 为取料总量
    merged: dict[int, dict[str, Any]] = {}
    for index, comp in matched.items():
        entry = merged.get(comp.id)
        if entry is None:
            merged[comp.id] = {
                "component": comp,
                "need": 0,
                "line_indexes": [],
                "raws": [],
                "_reported": False,
            }
            entry = merged[comp.id]
        entry["need"] += lines[index].quantity
        entry["line_indexes"].append(index)
        entry["raws"].append(lines[index].raw)

    missing: list[dict] = []
    for index, line in enumerate(lines):
        comp = matched.get(index)
        if comp is None:  # 找不到对应槽位
            if any(choice["line_index"] == index for choice in choices):
                continue
            missing.append({
                "reason": "not_found",
                "raw": line.raw,
                "name": line.name,
                "value": line.value,
                "package": line.package,
                "quantity": line.quantity,
                "component_id": None,
                "available": None,
            })
            continue
        entry = merged[comp.id]
        if entry["_reported"]:
            continue
        entry["_reported"] = True
        if entry["need"] > comp.quantity:  # 找到但量不足
            missing.append({
                "reason": "shortage",
                "raw": "；".join(entry["raws"]),
                "name": comp.name,
                "value": comp.value,
                "package": comp.package,
                "quantity": entry["need"],
                "component_id": comp.id,
                "available": comp.quantity,
            })

    # 足料者成步骤：先非 None（数值小在前），None 灯位排最后
    steps = [
        {"component": entry["component"], "quantity": entry["need"],
         "line_indexes": entry["line_indexes"]}
        for entry in merged.values()
        if entry["need"] <= entry["component"].quantity
    ]
    steps.sort(key=lambda item: (
        item["component"].led_index is None,
        item["component"].led_index if item["component"].led_index is not None
        else 0,
    ))

    return {
        "requested": sum(line.quantity for line in lines),
        "steps": steps,
        "missing": missing,
        "choices": choices,
    }


# ---------- Excel BOM（嘉立创/EasyEDA 导出等） ----------

# 表头别名 → 逻辑列；命中任一即认为该列承载对应信息
_HEADER_ALIASES = {
    "name": ("name", "part", "comment", "物料", "名称", "元件", "器件"),
    "value": ("value", "spec", "规格", "值", "容值"),
    "footprint": ("footprint", "package", "封装", "焊盘"),
    "quantity": ("quantity", "qty", "count", "数量", "pcs"),
    # 采购/识别依据：Manufacturer Part 与 Supplier Part（立创商城编号）
    "manufacturer_part": ("manufacturer part", "mpn", "mfr part",
                          "厂商料号", "料号", "型号"),
    "supplier_part": ("supplier part", "lcsc", "立创编号",
                      "商城编号", "购买编号", "供应商料号"),
}

# 封装归一：C0402 / R0603 / L0603 → 0402 / 0603（嘉立创习惯前缀字母+四位）
_PKG4_RE = re.compile(r"^[A-Za-z]{0,3}(\d{4})")

# 常见非阻容感元件家族指纹（型号前缀/关键词 → 中文族名）。
# 只做“展示兜底”，不改写用户原始型号（识别/购买仍看原始 Name 与 MPN）。
_KIND_PATTERNS: list[tuple[str, str]] = [
    (r"^(stm32|stm8|esp32|esp8266|atmega|attiny|atm|rp2040|ch32|gd32)",
     "芯片"),
    (r"^(w25q|at24|at25|mx25|gd25|sst25|flash)", "存储芯片"),
    (r"^(lm|ams|tl|tp|me|xc|rt|sg|mc34063|7805|1117|mp23|mp1\d)",
     "电源/模拟芯片"),
    (r"^(ss|bat54|1n\d|s1[bm]|us1|es1|fr107|her|sr5|mbr)", "二极管"),
    (r"^(s\d\d\d\d|ao\d|2n\d|bc\d|bd\d|a\d\d\d|irf|irl|csd)",
     "三极管/MOS"),
    (r"mhz|khz|晶振|crystal", "晶振"),
    (r"^(sw|k2|ts-|tl\d|skq|轻触|按键)", "开关"),
    (r"^(xh|ph|zh|pa|pb|pt|sh|jst|xh2|usb|type-c|dc\d|排针|连接器)",
     "连接器"),
    (r"^(f\d|保险丝|fuse|littelfuse)", "保险丝"),
    (r"^(led|贴片led|发光二极管)", "LED"),
    (r"^(蜂鸣器|buzzer|喇叭)", "蜂鸣器"),
]


def _cell_text(value: Any) -> str:
    """单元格 → 干净字符串（数值转 int 表示，避免 11.0）。"""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _footprint_package(footprint: str) -> Optional[str]:
    """从 Footprint 列提取四位封装号（C0402/R0603/L0603 → 0402/0603）。

    只认数字四位；LQFP-100、SW-SMD、CRYSTAL-SMD 等板级封装不塞进
    名称/值文本（会被裸数字误解析成阻值），保持元件名原样。
    """
    fp = _cell_text(footprint)
    if not fp:
        return None
    m = _PKG4_RE.match(fp)
    return m.group(1) if m else None


def _family_cn(value: Optional[str]) -> str:
    """元件描述 → 中文族名（Excel 无名字行/购买清单展示兜底）。

    先按值归一化认阻容感，再按型号指纹认常见非阻容感器件；
    返回的是展示标签，绝不改写用户的原始型号。
    """
    text = (value or "").strip()
    norm = normalize_value(text) if text else None
    if norm:
        family = norm[1]
        return {"F": "电容", "R": "电阻", "H": "电感"}.get(family, "元件")
    lower = text.lower()
    for pattern, kind in _KIND_PATTERNS:
        if re.search(pattern, lower):
            return kind
    return "元件"


def _compose_row(row: list[Any], mapping: dict) -> Optional[BomLine]:
    """分别读取表格字段，保留型号并以数量列为准。"""
    def cell(key: str) -> str:
        index = mapping.get(key)
        if index is None or index >= len(row):
            return ""
        return _cell_text(row[index])

    name = cell("name")
    value = cell("value")
    fp = cell("footprint")
    raw_qty = cell("quantity")
    mpn = cell("manufacturer_part")
    supplier = cell("supplier_part")

    if not any((name, value, fp, raw_qty, mpn, supplier)):
        return None
    pkg = _footprint_package(fp)
    try:
        quantity = int(raw_qty) if raw_qty else 1
    except ValueError as exc:
        raise ValueError(f"BOM 数量无效：{raw_qty}") from exc
    if quantity < 0:
        raise ValueError(f"BOM 数量无效：{raw_qty}")

    # 无独立值列时仅从完整标称值推断；裸数字型号不当成阻值。
    if not value and normalize_value(name) is not None:
        if not name.isdigit() or pkg:
            value = name
    display_name = name or mpn or supplier or _family_cn(value)
    if name == value and normalize_value(value) is not None:
        display_name = _family_cn(value)
    parts = [name]
    if value != name:
        parts.append(value)
    parts.extend([pkg or "", f"×{quantity}"])
    return BomLine(
        raw=" ".join(part for part in parts if part), name=display_name,
        value=value or None, package=pkg, quantity=quantity,
        manufacturer_part=mpn or None, supplier_part=supplier or None,
    )


def _find_header(rows: list[list[Any]]) -> Optional[tuple[int, dict]]:
    """在前几行找表头并返回行号与 {逻辑列: 列下标}。"""
    for row_index, row in enumerate(rows[:6]):
        if not row:
            continue
        cells = [_cell_text(c).lower() for c in row]
        mapping: dict = {}
        for key, aliases in _HEADER_ALIASES.items():
            for i, cell in enumerate(cells):
                if cell and cell in aliases:
                    mapping[key] = i
                    break
        # 至少要有名称类 + 数量/封装之一，才算表头
        if "name" in mapping and ("quantity" in mapping or "footprint" in mapping):
            return row_index, mapping
    return None


def _parse_table_rows(rows: list[list[Any]]) -> list[BomLine]:
    """解析带表头的二维行数据，供 xlsx 和 CSV 共用。"""
    header = _find_header(rows)
    if header is None:
        raise ValueError("未识别到表头（需要包含 名称/数量 或 名称/封装 的列）")
    header_index, mapping = header

    lines: list[BomLine] = []
    for row in rows[header_index + 1:]:
        if not any(_cell_text(cell) for cell in row):
            continue
        line = _compose_row(row, mapping)
        if line is not None:
            lines.append(line)
    return lines


def parse_csv_bytes(data: bytes) -> list[BomLine]:
    """解析 UTF-8 或 GB18030 编码的 BOM CSV 表格。"""
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = data.decode("gb18030")

    try:
        dialect = csv.Sniffer().sniff(text[:2048], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel

    rows = list(csv.reader(StringIO(text), dialect))
    return _parse_table_rows(rows)


def parse_excel_bytes(data: bytes) -> list[BomLine]:
    """解析 .xlsx BOM：识别常见表头并分别读取名称、封装与数量。"""
    try:
        workbook = load_workbook(BytesIO(data), data_only=True, read_only=True)
    except Exception as exc:
        raise ValueError(f"文件不是有效的 .xlsx：{exc}") from exc

    try:
        sheet = workbook.active
        rows = [list(row) for row in sheet.iter_rows(values_only=True)]
    finally:
        workbook.close()

    return _parse_table_rows(rows)
